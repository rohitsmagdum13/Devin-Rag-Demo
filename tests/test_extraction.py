from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas

from rag_demo.errors import AppError
from rag_demo.extraction import extract


def pdf_bytes(text: bool = True) -> bytes:
    buffer = BytesIO()
    pdf = canvas.Canvas(buffer)
    if text:
        pdf.drawString(72, 720, "First page policy")
    pdf.showPage()
    if text:
        pdf.drawString(72, 720, "Second page policy")
    pdf.save()
    return buffer.getvalue()


def test_pdf_pages() -> None:
    sections = extract("policy.PDF", pdf_bytes(), 100000)
    assert [section.page for section in sections] == [1, 2]
    assert "First page" in sections[0].text


def test_docx_paragraph_and_table() -> None:
    buffer = BytesIO()
    doc = Document()
    doc.add_paragraph("Annual policy")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "1200 allowance"
    doc.save(buffer)
    section = extract("policy.docx", buffer.getvalue(), 100000)[0]
    assert "Annual policy" in section.text and "1200 allowance" in section.text
    assert section.page is None


def test_text_bom() -> None:
    assert extract("policy.txt", b"\xef\xbb\xbfHello", 100)[0].text == "Hello"


@pytest.mark.parametrize(
    "name,content,match",
    [
        ("file.csv", b"value", "Unsupported"),
        ("empty.txt", b"", "empty"),
        ("blank.txt", b"  \n ", "No readable"),
        ("invalid.txt", b"\xff", "UTF-8"),
        ("binary.txt", b"hello\x00", "binary"),
        ("broken.pdf", b"not a pdf", "Invalid PDF"),
        ("broken.pdf", b"%PDF-invalid", "Could not read"),
        ("broken.docx", b"not a zip", "Could not read"),
    ],
)
def test_bad_files(name: str, content: bytes, match: str) -> None:
    with pytest.raises(AppError, match=match):
        extract(name, content, 100000)


def test_oversized() -> None:
    with pytest.raises(AppError, match="exceeds"):
        extract("large.txt", b"abcd", 3)


def test_blank_pdf() -> None:
    with pytest.raises(AppError, match="No readable"):
        extract("blank.pdf", pdf_bytes(False), 100000)


def test_encrypted_pdf() -> None:
    output = BytesIO()
    writer = PdfWriter()
    writer.append_pages_from_reader(PdfReader(BytesIO(pdf_bytes())))
    writer.encrypt("test-only")
    writer.write(output)
    with pytest.raises(AppError, match="Encrypted"):
        extract("locked.pdf", output.getvalue(), 100000)
