from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from pypdf import PdfReader

from rag_demo.errors import AppError
from rag_demo.models import TextSection


def extract(name: str, content: bytes, max_bytes: int) -> list[TextSection]:
    extension = Path(name).suffix.lower()
    if extension not in {".pdf", ".docx", ".txt"}:
        raise AppError("Unsupported file. Upload a PDF, DOCX, or UTF-8 TXT file.")
    if not content:
        raise AppError("The document is empty. Upload a file containing readable text.")
    if len(content) > max_bytes:
        raise AppError(f"File exceeds the {max_bytes // (1024 * 1024)} MB upload limit.")
    try:
        if extension == ".txt":
            text = content.decode("utf-8-sig")
            if "\x00" in text:
                raise AppError("The TXT file contains binary data. Export it as UTF-8 text.")
            sections = [TextSection(text)]
        elif extension == ".pdf":
            if not content.lstrip().startswith(b"%PDF-"):
                raise AppError("Invalid PDF. Re-export the document as a PDF and retry.")
            reader = PdfReader(BytesIO(content))
            if reader.is_encrypted:
                raise AppError("Encrypted PDF. Remove its password before uploading.")
            sections = [
                TextSection(page.extract_text() or "", i + 1) for i, page in enumerate(reader.pages)
            ]
        else:
            with ZipFile(BytesIO(content)) as archive:
                if sum(info.file_size for info in archive.infolist()) > 100 * 1024 * 1024:
                    raise AppError("DOCX expands beyond the 100 MB safety limit.")
                if "word/document.xml" not in archive.namelist():
                    raise AppError("Invalid DOCX. Export the document as DOCX and retry.")
            doc = Document(BytesIO(content))
            parts = [paragraph.text for paragraph in doc.paragraphs]
            parts.extend(
                " | ".join(cell.text for cell in row.cells)
                for table in doc.tables
                for row in table.rows
            )
            sections = [TextSection("\n".join(parts))]
    except UnicodeDecodeError:
        raise AppError("TXT must be UTF-8 encoded. Re-save it as UTF-8 and retry.") from None
    except AppError:
        raise
    except Exception:
        raise AppError("Could not read the document. Re-export it and upload again.") from None
    cleaned = [
        TextSection(section.text.strip(), section.page)
        for section in sections
        if section.text.strip()
    ]
    if not cleaned:
        raise AppError("No readable text found. Scanned PDFs require OCR before uploading.")
    return cleaned
